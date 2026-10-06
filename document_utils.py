import re
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u2013": "-", "\u2014": "-", "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"[ \t]+", " ", text).strip()

def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(54)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"LEGALEASE\n{doc_type.upper()}")
    run.bold = True
    run.font.size = Pt(16)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            continue
        p = doc.add_paragraph()
        if re.match(r"^(\d+[\.\)]|SECTION\s+\d+)", line, re.I):
            r = p.add_run(line)
            r.bold = True
        else:
            p.add_run(line)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated draft | Review by a qualified legal professional")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()

class LegalPDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = doc_type

    def header(self):
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, "LEGALEASE", align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_font("Helvetica", "B", 11)
        self.cell(0, 7, self.doc_type.upper(), align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(0, 10, "LegalEase - AI-generated draft - Page " + str(self.page_no()), align="C")

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalPDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(3)
            continue
        if re.match(r"^(\d+[\.\)]|SECTION\s+\d+)", line, re.I):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, line)
            pdf.set_font("Helvetica", size=11)
        else:
            pdf.multi_cell(0, 6, line)

    return bytes(pdf.output())
